#!/usr/bin/env python3
"""Build the Advanced Bazel testing/validation workshop deck."""

import argparse

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

_parser = argparse.ArgumentParser(description=__doc__)
_parser.add_argument("--template", required=True)
_parser.add_argument("--output", required=True)
_args = _parser.parse_args()

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x0F, 0x1F, 0x33)   # matches dk1 in the template
GREY = RGBColor(0x62, 0x6D, 0x7C)   # matches dk2 in the template

CODE_FONT = "Roboto Mono"
BODY_FONT = "Poppins"


def clear_deck(prs):
    sldIdLst = prs.slides._sldIdLst
    rIds = [sldId.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            for sldId in list(sldIdLst)]
    for sldId in list(sldIdLst):
        sldIdLst.remove(sldId)
    part = prs.part
    for rId in rIds:
        try:
            part.drop_rel(rId)
        except Exception:
            pass


def clear_tf(tf):
    txBody = tf._txBody
    for p in txBody.findall(f"{{{A_NS}}}p"):
        txBody.remove(p)
    etree.SubElement(txBody, f"{{{A_NS}}}p")


def set_run(run, *, size=None, bold=None, italic=None, color=None, font=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if font is not None:
        run.font.name = font


LEVEL_INDENT = {0: (457200, -342900), 1: (914400, -342900), 2: (1371600, -342900)}


def add_paragraph(tf, text, *, size=18, bold=False, italic=False, color=WHITE,
                  level=0, bullet=True, bullet_color=None,
                  align=None, first=False, font=BODY_FONT):
    if first and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.level = level
    if align is not None:
        p.alignment = align

    pPr = p._p.get_or_add_pPr()
    for tag in ("buClr", "buChar", "buAutoNum", "buNone", "buFont",
                "buSzPts", "buSzPct"):
        for el in pPr.findall(f"{{{A_NS}}}{tag}"):
            pPr.remove(el)

    marL, indent = LEVEL_INDENT.get(level, LEVEL_INDENT[0])
    pPr.set("marL", str(marL))
    pPr.set("indent", str(indent))

    if bullet:
        bc = bullet_color if bullet_color is not None else color
        buClr = etree.SubElement(pPr, f"{{{A_NS}}}buClr")
        srgb = etree.SubElement(buClr, f"{{{A_NS}}}srgbClr")
        srgb.set("val", "{:02X}{:02X}{:02X}".format(bc[0], bc[1], bc[2]))
        buFont = etree.SubElement(pPr, f"{{{A_NS}}}buFont")
        buFont.set("typeface", BODY_FONT)
        buChar = etree.SubElement(pPr, f"{{{A_NS}}}buChar")
        buChar.set("char", "●" if level == 0 else ("○" if level == 1 else "■"))
    else:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, f"{{{A_NS}}}buNone")

    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, italic=italic, color=color, font=font)
    return p


def set_title(slide, text, *, size=32):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            tf = ph.text_frame
            clear_tf(tf)
            add_paragraph(tf, text, size=size, bold=True, color=WHITE,
                          bullet=False, first=True)
            return


def add_bullets(slide, items, *, ph_idx=1, base_size=18, color=WHITE):
    ph = None
    for p in slide.placeholders:
        if p.placeholder_format.idx == ph_idx:
            ph = p
            break
    if ph is None:
        return
    tf = ph.text_frame
    tf.word_wrap = True
    clear_tf(tf)
    first = True
    for it in items:
        text, level = (it if isinstance(it, tuple) else (it, 0))
        size = base_size if level == 0 else max(12, base_size - 3)
        add_paragraph(tf, text, size=size, level=level, bullet=True,
                      color=color, first=first)
        first = False


def add_code_card(slide, lines, *, top=1135353, height=3780000, size=13):
    """White rounded card + monospaced text laid out per template slide 18."""
    card_left = 416775
    card_width = 8316000
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Emu(card_left), Emu(top), Emu(card_width), Emu(height),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.fill.background()
    prstGeom = card._element.spPr.prstGeom
    avLst = prstGeom.find(f"{{{A_NS}}}avLst")
    if avLst is None:
        avLst = etree.SubElement(prstGeom, f"{{{A_NS}}}avLst")
    gd = etree.SubElement(avLst, f"{{{A_NS}}}gd")
    gd.set("name", "adj")
    gd.set("fmla", "val 2121")

    text_left = card_left + 297825
    text_top = top + 233400
    text_w = card_width - 915300
    text_h = height - 466800
    tb = slide.shapes.add_textbox(Emu(text_left), Emu(text_top),
                                  Emu(text_w), Emu(text_h))
    tb.fill.background()
    tb.line.fill.background()
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    clear_tf(tf)
    first = True
    for line in lines:
        italic = False
        color = DARK
        if isinstance(line, tuple):
            line, kind = line
            if kind == "comment":
                italic = True
                color = GREY
        add_paragraph(tf, line, size=size, color=color, bullet=False,
                      italic=italic, font=CODE_FONT, first=first)
        first = False
    return card


_CURRENT_BRANCH = [None]


def set_branch(branch):
    """Set (or clear with None) the solution branch stamped on subsequent slides."""
    _CURRENT_BRANCH[0] = branch


def stamp_branch(slide):
    """Stamp 'solution branch: <name>' at the bottom of `slide`, if any is set."""
    branch = _CURRENT_BRANCH[0]
    if not branch:
        return
    tb = slide.shapes.add_textbox(Emu(311700), Emu(4900000),
                                  Emu(8520600), Emu(250000))
    tb.fill.background()
    tb.line.fill.background()
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0);  tf.margin_bottom = Emu(0)
    clear_tf(tf)
    add_paragraph(tf, f"solution branch: {branch}",
                  size=11, italic=True, color=WHITE, bullet=False,
                  align=PP_ALIGN.RIGHT, first=True, font=CODE_FONT)


def new_slide(prs, layout):
    """Add a slide and stamp the current solution branch on it."""
    s = prs.slides.add_slide(layout)
    stamp_branch(s)
    return s


def section_slide(prs, title, *, subtitle=None, solution_branch=None):
    """Add a section-header slide. Section headers always reset the current
    branch: pass solution_branch for exercise sections, omit it for
    interludes to clear the stamp on subsequent slides."""
    set_branch(solution_branch)
    s = prs.slides.add_slide(L_SECTION)
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 0:
            tf = ph.text_frame
            clear_tf(tf)
            add_paragraph(tf, title, size=48, bold=True, color=WHITE,
                          bullet=False, align=PP_ALIGN.CENTER, first=True)
            if subtitle:
                add_paragraph(tf, subtitle, size=22, color=WHITE,
                              bullet=False, align=PP_ALIGN.CENTER)
    stamp_branch(s)
    return s


def strip_handle(slide):
    for sh in list(slide.shapes):
        if sh.has_text_frame and "@handle" in sh.text_frame.text:
            sh._element.getparent().remove(sh._element)


# ------------------------------- Build deck -------------------------------

prs = Presentation(_args.template)
clear_deck(prs)

L_TITLE = prs.slide_layouts[0]        # TITLE
L_SECTION = prs.slide_layouts[3]      # SECTION_HEADER
L_BODY = prs.slide_layouts[4]         # TITLE_AND_BODY
L_TITLE_ONLY = prs.slide_layouts[6]   # TITLE_ONLY


# ---------- Slide: Title ----------
s = prs.slides.add_slide(L_TITLE)
strip_handle(s)
for ph in s.placeholders:
    if ph.placeholder_format.idx == 0:
        tf = ph.text_frame
        clear_tf(tf)
        add_paragraph(tf, "Advanced Bazel", size=44, bold=True,
                      color=WHITE, bullet=False, first=True)
        add_paragraph(tf, "Rules Authoring, Performance Tuning and Remote Execution",
                      size=24, color=WHITE, bullet=False)
    elif ph.placeholder_format.idx == 1:
        tf = ph.text_frame
        clear_tf(tf)
        add_paragraph(tf, "Antonio Di Stefano  ·  Yannic Bonenberger",
                      size=18, bold=True, color=WHITE, bullet=False, first=True)


# ---------- Slide: Meet your speakers ----------
s = prs.slides.add_slide(L_BODY)
set_title(s, "Meet your speakers")
add_bullets(s, [
    "Antonio Di Stefano — EngFlow",
    ("Software Engineer", 1),
    "Yannic Bonenberger — EngFlow",
    ("Software Engineer", 1),
], base_size=20)


# ---------- Section: Testing & Validation ----------
section_slide(prs, "Testing & Validation")


# ---------- Slide: Four building blocks ----------
s = new_slide(prs, L_BODY)
set_title(s, "Four Bazel building blocks")
add_bullets(s, [
    "Test sharding — one test target, N parallel workers",
    "Validation actions — implicit checks via the _validation output group",
    "Tree artifacts — directory outputs whose contents are only known at execution time",
    "ctx.actions.map_directory — fanning out N actions over a tree artifact",
], base_size=20)


# ---------- Slide: What you start with ----------
s = new_slide(prs, L_BODY)
set_title(s, "What you start with")
add_bullets(s, [
    "A working minilang CLI at //cli:cli",
    "An example minilang program at //:example.mini",
    "A minimal mini_test rule that invokes 'minilang test <src>' and honors the exit code",
    ("bazel test //:example_test already passes", 1),
], base_size=18)


# ---------- Slide: minilang test — flags ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "minilang test")
add_code_card(s, [
    "$ bazel run //cli:cli -- test example.mini",
    "$ bazel run //cli:cli -- test example.mini \\",
    "    --shard-count 3 --shard-index 0",
    "$ bazel run //cli:cli -- test example.mini \\",
    "    --junit-xml out.xml",
    "",
    ("# --shard-count N   total shard count (round robin)", "comment"),
    ("# --shard-index i   0-based shard index; runs tests where i % N == index", "comment"),
    ("# --junit-xml PATH  JUnit-compatible XML report at PATH", "comment"),
], size=13)


# ---------- Exercise section: Test sharding ----------
section_slide(prs, "Test sharding",
              solution_branch="testing_and_validation_stage_sharding")


# ---------- Slide: sharding goal ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "Goal")
add_code_card(s, [
    'mini_test(',
    '    name = "example_test_sharded",',
    '    src  = "example.mini",',
    '    shard_count = 3,',
    ')',
    "",
    ("# bazel test //:example_test_sharded", "comment"),
    ("# → three parallel actions, one test case each", "comment"),
], size=14)


# ---------- Slide: Bazel's shard contract ----------
s = new_slide(prs, L_BODY)
set_title(s, "Bazel's contract with sharded tests")
add_bullets(s, [
    "TEST_TOTAL_SHARDS — number of shards. 1 means unsharded",
    "TEST_SHARD_INDEX — 0-based index of this shard",
    "TEST_SHARD_STATUS_FILE — the runner MUST create the file or update its mtime",
    ("Otherwise Bazel fails the test", 1),
    "shard_count is a built-in test attribute — you don't declare it on your rule",
    ("Your job: plumb env vars → CLI's --shard-count / --shard-index", 1),
], base_size=17)


# ---------- Section: Mutation testing ----------
section_slide(prs, "Mutation testing",
              subtitle="and why the naïve approach is wrong")


# ---------- Slide: What is mutation testing ----------
s = new_slide(prs, L_BODY)
set_title(s, "Mutation testing in one paragraph")
add_bullets(s, [
    "Grade a test suite by perturbing the program under test",
    ("Flip < to <=, swap + for -, replace true with false, …", 1),
    "Re-run the tests",
    ("A test fails → mutation is killed (good)", 1),
    ("Every test passes → mutation survived (bad — tests are blind to it)", 1),
], base_size=19)


# ---------- Slide: mutate subcommands ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "minilang mutate — subcommands")
add_code_card(s, [
    "$ bazel run //cli:cli -- mutate list example.mini -o mutations/",
    ("# mutations/mutations.json + mNNNN/mutation.json per candidate", "comment"),
    "",
    "$ bazel run //cli:cli -- mutate patch \\",
    "    --mutation mNNNN/mutation.json \\",
    "    --source example.mini \\",
    "    --output mutated.mini",
    "",
    "$ bazel run //cli:cli -- mutate test mutated.mini \\",
    "    --junit-xml test.xml \\",
    "    --mutation-json mNNNN/mutation.json",
    ("# inverted exit code: 0 iff at least one test failed", "comment"),
], size=12)


# ---------- Slide: The tempting-but-wrong design ----------
s = new_slide(prs, L_BODY)
set_title(s, "The tempting-but-wrong design")
add_bullets(s, [
    "Add a 'mutation_test' mode to mini_test that loops over every mutation at test time",
    "Don't do this. Concretely:",
    ("One action = one worker — no parallelism, 24 mutations run sequentially", 1),
    ("All-or-nothing caching — one line changes, everything reruns", 1),
    ("All-or-nothing failure — one survivor = red test, others invisible", 1),
    ("All-or-nothing logs — everything smashed into one test.log", 1),
], base_size=17)


# ---------- Slide: three phases ----------
s = new_slide(prs, L_BODY)
set_title(s, "The right shape: three phases")
add_bullets(s, [
    "Enumeration — what mutations exist? One action, produces a manifest",
    "Mutation — apply mutation m0007. One action per mutation",
    "Verification — does m0007 survive? One action per mutation",
    "Bazel schedules, caches, and reports each mutation's outcome independently",
], base_size=19)


# ---------- Exercise section: Validation actions ----------
section_slide(prs, "Validation actions",
              solution_branch="testing_and_validation_stage_validation")


# ---------- Slide: Aspect goal ----------
s = new_slide(prs, L_BODY)
set_title(s, "Goal")
add_bullets(s, [
    "Create mini_mutation_testing_aspect",
    "Attaches an extra action to every mini_test target",
    ("For now the work is trivial (a placeholder file)", 1),
    ("The point is the wiring, not what the aspect does yet", 1),
], base_size=19)


# ---------- Slide: First cut — explicit output group ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "First cut: an explicit output group")
add_code_card(s, [
    "return [OutputGroupInfo(mutations = depset([marker]))]",
    "",
    "$ bazel build \\",
    "    --aspects=//:defs.bzl%mini_mutation_testing_aspect \\",
    "    --output_groups=mutations \\",
    "    //:example_test",
    "",
    ("# Forget --output_groups and the aspect's action never runs.", "comment"),
    ("# Every consumer has to know the group's name. Awkward.", "comment"),
], size=13)


# ---------- Slide: The _validation magic key ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "The fix: _validation")
add_code_card(s, [
    "return [OutputGroupInfo(_validation = depset([marker]))]",
    "",
    "$ bazel build \\",
    "    --aspects=//:defs.bzl%mini_mutation_testing_aspect \\",
    "    //:example_test",
    ("# No --output_groups. Bazel runs it automatically.", "comment"),
    "",
    "$ bazel build --aspects=... \\",
    "    --run_validations=false //:example_test",
    ("# Escape hatch for rapid iteration; --keep_going shows every failure.", "comment"),
], size=13)


# ---------- Exercise section: Tree artifacts ----------
section_slide(prs, "Tree artifacts",
              solution_branch="testing_and_validation_stage_tree_artifacts")


# ---------- Slide: Tree-artifacts goal ----------
s = new_slide(prs, L_BODY)
set_title(s, "Goal")
add_bullets(s, [
    "Replace the placeholder work in the aspect",
    "Have it run 'minilang mutate list <src> -o <dir>'",
    "Expose <dir> as a directory output — a tree artifact",
], base_size=20)


# ---------- Slide: Tree-artifacts building blocks ----------
s = new_slide(prs, L_BODY)
set_title(s, "Building blocks")
add_bullets(s, [
    "Tree artifact: directory output known only at execution time",
    ("Single artifact in the action graph; downstream actions can enumerate its entries", 1),
    "ctx.actions.declare_directory(name)",
    ("Returns a File; file.path is the directory the action must populate", 1),
    "ctx.actions.run(executable, inputs, outputs, arguments, …)",
    ("The producing action lists the tree artifact in outputs", 1),
    "ctx.actions.args() — assemble the command line without string concatenation",
], base_size=16)


# ---------- Slide: Tree-artifacts verify ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "Verify")
add_code_card(s, [
    "$ bazel build --aspects=... //:example_test",
    "$ ls bazel-bin/example_test.mutations/",
    "mutations.json",
    "m0000/  m0001/  m0002/  ...",
    "",
    ("# The count depends on the source — Bazel didn't know at analysis time.", "comment"),
    ("# That's the whole point of tree artifacts.", "comment"),
], size=14)


# ---------- Exercise section: map_directory ----------
section_slide(prs, "ctx.actions.map_directory",
              solution_branch="testing_and_validation_stage_map_directory")


# ---------- Slide: map_directory goal ----------
s = new_slide(prs, L_BODY)
set_title(s, "Goal — two actions per mutation")
add_bullets(s, [
    "For each mutation subdirectory produced above:",
    ("mutate patch → mutated source in the results tree", 1),
    ("mutate test  → runs the test suite (inverted exit code)", 1),
    "A mutation survives iff its test action exits non-zero",
    ("Bazel surfaces the one-liner: 'survived: <id> <op> at <file>:<line>:<col>'", 1),
], base_size=18)


# ---------- Slide: map_directory signature ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "ctx.actions.map_directory")
add_code_card(s, [
    "ctx.actions.map_directory(",
    "    input_directories,     # tree artifact File",
    "    output_directories,    # tree artifact File",
    "    tools,                 # FilesToRunProvider",
    "    additional_inputs = {}, additional_params = {},",
    "    implementation,        # top-level function",
    "    mnemonic = None,",
    ")",
    "",
    ("# Declared once at analysis time.", "comment"),
    ("# Callback creates N distinct actions from execution-time data.", "comment"),
    ("# Each action is independently cacheable and failable.", "comment"),
], size=13)


# ---------- Slide: callback signature ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "The implementation callback")
add_code_card(s, [
    "def impl(template_ctx, *, input_directories,",
    "         output_directories, additional_inputs,",
    "         tools, additional_params):",
    "    ...",
    "",
    ("# input_directories[key] → ExpandedDirectory", "comment"),
    ("#   .directory, .children (each has .tree_relative_path, .path)", "comment"),
    ("# output_directories[key] → tree artifact File declared outside", "comment"),
    ("# tools[key] → FilesToRunProvider (use .executable)", "comment"),
    ("# Callback must be top-level (no lambdas, no nested defs)", "comment"),
], size=13)


# ---------- Slide: template_ctx ----------
s = new_slide(prs, L_BODY)
set_title(s, "template_ctx — restricted surface")
add_bullets(s, [
    "template_ctx.args() — same as ctx.actions.args()",
    "template_ctx.declare_file(filename, *, directory)",
    ("Declares a file inside the given output tree artifact", 1),
    "template_ctx.declare_subdirectory(name, *, directory)",
    "template_ctx.run(outputs, executable, inputs, tools, arguments, progress_message)",
    ("No mnemonic. No run_shell. Outputs must be declared via declare_file", 1),
    "Every read → in inputs=[…]. Every write → declared + in outputs=[…]",
], base_size=16)


# ---------- Slide: map_directory verify ----------
s = new_slide(prs, L_TITLE_ONLY)
set_title(s, "Verify")
add_code_card(s, [
    "$ bazel build --aspects=... --keep_going //:example_test",
    "",
    ("# ≈ 2 × N actions (patch + test per mutation)", "comment"),
    ("# Surviving mutations report as failed actions:", "comment"),
    "",
    "  survived: m0006 *->/ at mutated.mini:15:19",
], size=13)


# ---------- Slide: Takeaways ----------
s = new_slide(prs, L_BODY)
set_title(s, "Takeaways")
add_bullets(s, [
    "Test sharding is a rule-author's job: forward env vars, touch the status file",
    "_validation turns extra checks into implicit build-time work",
    "Tree artifacts let you produce work whose shape is decided at execution time",
    "map_directory bridges analysis and execution — N actions from execution-time data",
    "Pattern generalizes: produce a manifest → fan out per-entry → surface as validation",
    ("Code gen, per-package linting, per-file coverage, ABI diffs, …", 1),
], base_size=16)


# ---------- Slide: Q&A ----------
s = prs.slides.add_slide(L_SECTION)
for ph in s.placeholders:
    if ph.placeholder_format.idx == 0:
        tf = ph.text_frame
        clear_tf(tf)
        add_paragraph(tf, "Q&A", size=88, bold=True, color=WHITE,
                      bullet=False, align=PP_ALIGN.CENTER, first=True)


prs.save(_args.output)
print(f"Wrote {_args.output}")
print(f"Slide count: {len(prs.slides)}")
